"""Safe text extraction from PDF / DOCX resumes."""

import zipfile


class ExtractionError(Exception):
    """Raised when a resume cannot be safely read."""
    pass


MAX_PDF_PAGES = 10
MAX_DOCX_UNCOMPRESSED_SIZE = 60 * 1024 * 1024


def normalize_extension(ext: str) -> str:
    """
    Normalize a file extension.

    Examples:
        "PDF"  -> ".pdf"
        ".PDF" -> ".pdf"
        "docx" -> ".docx"
    """
    ext = (ext or "").strip().lower()

    if not ext:
        return ""

    if not ext.startswith("."):
        ext = "." + ext

    return ext


def looks_like(ext: str, head: bytes) -> bool:
    """
    Check file magic bytes so a renamed file cannot masquerade
    as a PDF or DOCX.
    """

    ext = normalize_extension(ext)

    if ext == ".pdf":
        return head.startswith(b"%PDF")

    if ext == ".docx":
        # DOCX is a ZIP-based format.
        return head.startswith(b"PK\x03\x04")

    return False


def extract_text(path, ext: str) -> str:
    """
    Extract text from a PDF or DOCX file.

    Returns:
        str: Extracted resume text.

    Raises:
        ExtractionError: If the file cannot be read.
    """

    ext = normalize_extension(ext)

    if ext == ".pdf":
        return _pdf(path)

    if ext == ".docx":
        return _docx(path)

    raise ExtractionError(
        "Unsupported file type. Please upload a PDF or DOCX file."
    )


def _pdf(path) -> str:
    """Extract text from a PDF."""

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError(
            "PDF support is not installed. Install pypdf."
        ) from exc

    try:
        reader = PdfReader(path)
    except Exception as exc:
        raise ExtractionError(
            "This PDF could not be opened. "
            "It may be corrupted or invalid."
        ) from exc

    # Check password protection.
    if reader.is_encrypted:
        raise ExtractionError(
            "This PDF is password-protected. "
            "Remove the password and upload it again."
        )

    # Check for empty PDF.
    try:
        page_count = len(reader.pages)
    except Exception as exc:
        raise ExtractionError(
            "Could not read the pages in this PDF."
        ) from exc

    if page_count == 0:
        raise ExtractionError(
            "This PDF doesn't contain any pages."
        )

    extracted_pages = []

    # Only process the first MAX_PDF_PAGES pages.
    for page_number, page in enumerate(
        reader.pages[:MAX_PDF_PAGES],
        start=1,
    ):
        try:
            text = page.extract_text() or ""

            if text.strip():
                extracted_pages.append(text.strip())

        except Exception as exc:
            # Don't immediately fail because one page is problematic.
            print(
                f"Warning: Could not extract page {page_number}: "
                f"{repr(exc)}"
            )
            continue

    text = "\n\n".join(extracted_pages).strip()

    # Important:
    # A valid PDF can contain only images/scans and therefore
    # have no extractable text.
    if not text:
        raise ExtractionError(
            "We couldn't extract text from this PDF. "
            "The PDF may be scanned or image-based. "
            "Please upload a PDF containing selectable text."
        )

    return text


def _docx(path) -> str:
    """Extract text from a DOCX file."""

    try:
        import docx
    except ImportError as exc:
        raise ExtractionError(
            "DOCX support is not installed. Install python-docx."
        ) from exc

    # ---------------------------------------------------------
    # ZIP validation / zip-bomb protection
    # ---------------------------------------------------------

    try:
        with zipfile.ZipFile(path, "r") as archive:

            if not archive.testzip() is None:
                raise ExtractionError(
                    "This DOCX file appears to be corrupted."
                )

            total_uncompressed_size = sum(
                info.file_size
                for info in archive.infolist()
            )

            if total_uncompressed_size > MAX_DOCX_UNCOMPRESSED_SIZE:
                raise ExtractionError(
                    "This document is too large to process."
                )

    except ExtractionError:
        raise

    except zipfile.BadZipFile as exc:
        raise ExtractionError(
            "This DOCX file is corrupted or invalid."
        ) from exc

    except Exception as exc:
        raise ExtractionError(
            "This DOCX file could not be inspected."
        ) from exc

    # ---------------------------------------------------------
    # Open DOCX
    # ---------------------------------------------------------

    try:
        document = docx.Document(path)

    except Exception as exc:
        raise ExtractionError(
            "This DOCX file could not be opened. "
            "It may be corrupted or invalid."
        ) from exc

    lines = []

    # ---------------------------------------------------------
    # Paragraphs
    # ---------------------------------------------------------

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            lines.append(text)

    # ---------------------------------------------------------
    # Tables
    # ---------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            cells = []

            for cell in row.cells:

                text = cell.text.strip()

                if text:
                    cells.append(text)

            if cells:
                lines.append(" | ".join(cells))

    # ---------------------------------------------------------
    # Final text
    # ---------------------------------------------------------

    text = "\n".join(lines).strip()

    if not text:
        raise ExtractionError(
            "We couldn't extract any text from this DOCX file."
        )

    return text